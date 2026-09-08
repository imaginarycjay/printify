import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

output_dir = "/home/imaginarycjay/printify/capstone_paper/capstone_draft_folder/outline_draft"
os.makedirs(output_dir, exist_ok=True)
docx_path = os.path.join(output_dir, "Chapter_1_Introduction_Draft_v2.docx")
md_path = os.path.join(output_dir, "Chapter_1_Introduction_Draft_v2.md")

doc = docx.Document()

# Page setup: Letter size, USM standard margins: Left 1.5", Top 1.0", Right 1.0", Bottom 1.0"
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11.0)
section.left_margin = Inches(1.5)
section.right_margin = Inches(1.0)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)

# Configure default font to Arial 12pt
style_normal = doc.styles["Normal"]
font_normal = style_normal.font
font_normal.name = "Arial"
font_normal.size = Pt(12)
font_normal.color.rgb = RGBColor(0, 0, 0)
style_normal.paragraph_format.line_spacing = 1.5
style_normal.paragraph_format.space_after = Pt(6)
style_normal.paragraph_format.space_before = Pt(0)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    return p

def add_chapter_heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    return p

def add_section_heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    return p

def add_subsection_heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    return p

def add_body_paragraph(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    return p

def add_bullet_item(bold_prefix, text):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Arial"
        r_bold.font.size = Pt(12)
        r_bold.font.bold = True
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(12)
    return p

def add_numbered_item(num_str, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    r_num = p.add_run(num_str + " ")
    r_num.font.name = "Arial"
    r_num.font.size = Pt(12)
    r_num.font.bold = True
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(12)
    return p

def add_definition_item(term, definition):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    r_term = p.add_run(term)
    r_term.font.name = "Arial"
    r_term.font.size = Pt(12)
    r_term.font.bold = True
    r_def = p.add_run(" – " + definition)
    r_def.font.name = "Arial"
    r_def.font.size = Pt(12)
    return p

# --- BUILD DOCUMENT CONTENT ---

# Header Information
add_title("INTEGRATED DYNAMIC ORDER, JOB SCHEDULING, AND INVENTORY MANAGEMENT SYSTEM FOR PRINTING SERVICES WITH AUTOMATED REPLENISHMENT")

p_authors = doc.add_paragraph()
p_authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_authors.paragraph_format.space_after = Pt(18)
p_authors.paragraph_format.line_spacing = 1.5
r1 = p_authors.add_run("KAYE P. COMISSION\nCHRISTIAN JAMES J. PEREZ\n\nBACHELOR OF SCIENCE IN INFORMATION SYSTEMS\n\nAUGUST 2026")
r1.font.name = "Arial"
r1.font.size = Pt(12)
r1.font.bold = True

# Chapter Title
add_chapter_heading("INTRODUCTION")

# Background paragraphs
add_body_paragraph("The printing and binding industry is a fast-paced and highly dynamic sector that caters to a diverse range of client needs, from simple document printing and seasonal thesis binding to large-scale marketing materials. As educational institutions, corporate entities, and local businesses continuously generate demand, printing service providers are heavily challenged to maintain operational efficiency. Managing custom orders requires precision not only in recording customer specifications but also in allocating resources, tracking production progress, and ensuring that raw materials are always available. However, many local printing shops still rely on traditional, manual methods or disjointed software that cannot adapt to the varying workflows of different printing services.")

add_body_paragraph("Many local printing establishments experience significant operational bottlenecks, especially during peak seasons such as the end of academic semesters or local events. Their current manual setups struggle to seamlessly bridge the gap between order intake, production scheduling, and inventory tracking. When a sudden influx of orders occurs, tracking which employee or machine is assigned to a specific job becomes chaotic. Furthermore, the lack of real-time inventory monitoring often leads to unexpected stockouts of crucial materials—such as specific paper types, ink, and binding covers—resulting in delayed orders, loss of potential revenue, and decreased customer satisfaction.")

add_body_paragraph("To address these interconnected challenges, there is a vital need for a system that is not only automated but also highly adaptable to the specific services a shop offers. Traditional management systems are often rigidly hardcoded, preventing business owners from easily adding new product variations or pricing formulas as their business scales. Furthermore, without a data-driven approach to inventory, business owners are forced to guess their restocking schedules based on intuition rather than actual material consumption.")

add_body_paragraph("This project aims to design and develop an Integrated Dynamic Order, Job Scheduling, and Inventory Management System for Printing Services with Automated Replenishment. The core innovation of this system lies in its configurable architecture, which acts as a dynamic platform allowing administrators to define, add, or modify their own printing services, pricing variables, and material requirements without altering the source code. To optimize production, a dedicated Job Scheduling module will visually track the progress of each order and allocate labor and machinery efficiently. Finally, to eliminate the problem of material shortages, the system features an Automated Replenishment module. Instead of relying on long-term predictive models that require years of historical data, this module calculates the real-time \"burn rate\" of raw materials based on active sales volume and automatically alerts the management when stocks reach a critical reorder point.")

add_body_paragraph("By integrating these features into a single, cohesive web-based platform, this system intends to modernize the operational workflow of printing establishments, ensuring that orders are fulfilled on time, production is systematically scheduled, and inventory is optimally maintained.")

# Significance of the Project
add_section_heading("Significance of the Project")

add_body_paragraph("The realization of this project will profoundly benefit various stakeholders within the printing and binding industry by transforming disjointed manual processes into a streamlined, automated workflow.")

add_bullet_item("Business Owners and Management – ", "They will gain a dynamic and highly adaptable platform that allows them to effortlessly configure, add, or modify printing services and pricing structures without needing complex source code alterations. Furthermore, the automated inventory replenishment module will eliminate the guesswork in restocking, calculating real-time material burn rates to prevent costly stockouts and maximize potential revenue, especially during peak seasons.")

add_bullet_item("Employees and Production Staff – ", "The system will significantly reduce the chaos associated with sudden influxes of orders. The dedicated job scheduling module will provide a clear, visual queue of assigned tasks and deadlines, preventing overlapping duties and allowing staff to focus efficiently on production rather than manual coordination.")

add_bullet_item("Customers and Clients – ", "They will experience highly reliable service. With optimized order tracking, systematic job delegation, and guaranteed raw material availability, clients will benefit from faster turnaround times, strict adherence to deadlines, and an overall enhanced customer experience.")

add_bullet_item("Future Researchers and Developers – ", "This study will serve as a substantial foundation for future researchers. The integration of a dynamic, configurable service architecture combined with a short-term, data-driven material replenishment logic provides a modern software framework. Future IT and Information Systems scholars can utilize this project as a baseline reference for designing scalable, industry-specific management systems that do not rely on long-term, data-heavy predictive machine learning models.")

# Statement of the Problem
add_section_heading("Statement of the Problem")

add_body_paragraph("This study addresses the overarching problem of how to design and develop an Integrated Dynamic Order, Job Scheduling, and Inventory Management System for Printing Services with Automated Replenishment to effectively resolve the manual operational bottlenecks experienced by printing establishments. Currently, the reliance on disjointed, traditional processes creates significant frustrations for all primary stakeholders. For customers, the absence of real-time order tracking leads to the inconvenience of repeatedly visiting or contacting the establishment just to check the status of their print jobs, often resulting in dissatisfaction when orders are unexpectedly delayed. For production staff and employees, depending on manual, paper-based queuing systems creates a chaotic workflow where physical job tickets are easily misplaced, miscommunicated, or improperly prioritized, leading to overlapping duties and unfulfilled deadlines. Meanwhile, business owners and administrators are hindered by rigid software that cannot adapt to new services, and they are forced to rely on pure guesswork and intuition to manage inventory, which inevitably leads to unexpected material stockouts and lost revenue during peak seasons.")

add_body_paragraph("To resolve these interconnected issues, the study must explicitly answer how to formulate a dynamic order management architecture that allows administrators to seamlessly configure services and pricing on the fly, eliminating rigid system constraints. Furthermore, it seeks to determine how to develop a job scheduling module that digitizes the production queue to visually track order progress and allocate labor, thereby eliminating lost paperwork and staff confusion. Concurrently, it addresses how to implement an automated inventory replenishment module that calculates real-time material burn rates to replace manual inventory guesswork with data-driven reorder alerts. Finally, the study must answer how to integrate these targeted solutions into a single cohesive platform, and how to evaluate its overall acceptability and performance utilizing the ISO 25010 standard for functional suitability, usability, performance efficiency, and security.")

# Objectives of the Project
add_section_heading("Objectives of the Project")

add_body_paragraph("The general objective of the project is to design and develop an Integrated Dynamic Order, Job Scheduling, and Inventory Management System for Printing Services with Automated Replenishment to modernize and automate the operational workflows of printing establishments, from customizable order intake and production tracking to data-driven inventory maintenance.")

p_spec = doc.add_paragraph()
p_spec.paragraph_format.first_line_indent = Inches(0.5)
p_spec.paragraph_format.line_spacing = 1.5
p_spec.paragraph_format.space_after = Pt(4)
r_spec = p_spec.add_run("Specifically, this study aims to:")
r_spec.font.name = "Arial"
r_spec.font.size = Pt(12)

add_numbered_item("1.", "formulate a dynamic order management architecture that enables administrators to configure, add, or modify printing services, pricing variables, and material requirements without altering the source code;")
add_numbered_item("2.", "develop a job scheduling module that digitizes the production queue, visually tracks order progress, and systematically allocates labor and machinery to ensure timely fulfillment;")
add_numbered_item("3.", "implement an automated inventory replenishment module that calculates the real-time burn rate of materials based on active sales volume and triggers alerts at critical reorder points;")
add_numbered_item("4.", "integrate the dynamic ordering, job scheduling, and inventory replenishment modules into a single, cohesive web-based platform; and")
add_numbered_item("5.", "evaluate the overall acceptability and performance of the developed system using the ISO 25010 software quality standard, specifically focusing on functional suitability, usability, performance efficiency, and security.")

# Expected Outputs of the Project
add_section_heading("Expected Outputs of the Project")

add_body_paragraph("The expected outputs of the project will include the following:")

add_numbered_item("1.", "A fully functional web-based platform that integrates dynamic order configuration, job scheduling, and inventory replenishment tailored specifically for printing and binding service operations.")
add_numbered_item("2.", "A dynamic order management module that grants business administrators the flexibility to seamlessly create, customize, and modify various printing services and pricing structures without needing backend code alterations.")
add_numbered_item("3.", "A customer-facing portal that provides an intuitive interface for clients to easily submit their specific printing requirements and track the real-time progress of their orders, eliminating the need for frequent physical follow-ups.")
add_numbered_item("4.", "A digital job scheduling module that equips production staff with a visual, organized queue of pending and ongoing tasks, allowing for the efficient allocation of labor and machinery to meet strict fulfillment deadlines.")
add_numbered_item("5.", "An automated inventory replenishment module that continuously tracks material consumption and calculates real-time burn rates based on active sales, providing management with data-driven alerts before essential printing supplies reach critical depletion levels.")
add_numbered_item("6.", "A comprehensive system documentation and evaluation report detailing the software's performance, functional suitability, usability, and security based on the ISO 25010 quality standards.")

# Limitations of the Project
add_section_heading("Limitations of the Project")

add_body_paragraph("The scope of this project is strictly bounded to the order configuration, job scheduling, and inventory replenishment of printing services. The system is not designed to function as a full-scale Enterprise Resource Planning (ERP) software; thus, it completely excludes comprehensive financial accounting, third-party logistics or courier integration, and Human Resource modules such as payroll and biometric employee attendance. Furthermore, the system relies on manual payment verification (e.g., uploading of GCash receipts) and does not integrate with automated bank APIs or live payment gateways. Lastly, the software is purely web-based and does not feature Internet of Things (IoT) capabilities; it will not directly interface with physical printing or binding machines to automatically read hardware statuses or ink levels.")

add_subsection_heading("People")
add_body_paragraph("The system is designed to facilitate seamless interaction among three primary users:")
add_bullet_item("Administrators / Business Owners: ", "They have full access to the dynamic configuration of the system. They are responsible for setting up the types of printing services offered, managing pricing structures, monitoring the real-time inventory burn rate, and verifying customer payments.")
add_bullet_item("Production Staff / Employees: ", "They utilize the job scheduling module to view their assigned tasks, update the production progress of each order, and systematically deduct raw materials from the digital inventory as jobs are completed.")
add_bullet_item("Customers / Clients: ", "They access the customer-facing web portal to easily browse offered printing services, submit specific order requirements, upload necessary files, and track the real-time status of their print jobs without needing physical follow-ups.")

add_subsection_heading("Data")
add_body_paragraph("The system will process, store, and manage essential operational data to ensure smooth transactions. This includes customer profiles and contact details, transaction histories, and uploaded print files. On the administrative side, the database will handle the dynamic service configurations (custom pricing and variations), real-time job order statuses, and automated inventory records, specifically the current stock levels, material burn rates, and reorder point thresholds.")

add_subsection_heading("Process")
add_body_paragraph("The system automates the core operational workflow of the printing establishment. The process begins with dynamic order intake, where customers select services configured by the admin. It proceeds to manual payment verification, followed by the automated routing of the confirmed order into the digital job scheduling queue. During production, the system systematically allocates tasks to specific staff members. As production progresses, the system processes inventory deductions based on material consumption, ultimately triggering automated restocking alerts when supplies reach critical levels.")

add_subsection_heading("Hardware")
add_body_paragraph("To effectively run and access the web-based system, users only require standard web-capable devices such as desktop computers, laptops, tablets, or smartphones. A stable internet connection is required for all actors (administrators, staff, and customers) to access the cloud-hosted platform and synchronize data in real-time.")

# Operational Definition of Terms
add_section_heading("Operational Definition of Terms")

add_body_paragraph("To ensure clarity and establish a shared understanding of the concepts discussed throughout this study, the following terms are operationally defined. These definitions reflect how each technical and business-specific term is explicitly utilized within the context of the developed system's dynamic architecture, modules, and workflows, rather than their general dictionary meanings:")

add_definition_item("Administrator", "Refers to the business owner or authorized management personnel who has full access to the system. They are responsible for dynamically configuring printing services, managing pricing variables, monitoring the inventory burn rate, and verifying customer payments.")
add_definition_item("Automated Replenishment", "Refers to the specific system module that eliminates manual stock guesswork by calculating real-time material consumption and automatically triggering alerts when supplies reach a critical level.")
add_definition_item("Burn Rate", "In this study, it refers to the calculated velocity or speed at which specific raw printing materials (e.g., paper, ink, binding covers) are consumed based on active, short-term sales volume.")
add_definition_item("Customer", "Refers to the clients or end-users who utilize the customer-facing web portal to browse offered printing services, submit specific order requirements, upload files, and visually track the real-time status of their print jobs.")
add_definition_item("Dynamic Configuration", "Refers to the core architectural feature of the system that allows administrators to seamlessly create, add, or modify printing services and pricing structures through the user interface without requiring continuous modifications to the backend source code.")
add_definition_item("Inventory Management", "Refers to the digital tracking system used to monitor the availability, quantity, and depletion of raw materials required for production, which automatically updates as production staff complete their assigned tasks.")
add_definition_item("Job Scheduling", "Refers to the digital production queue that replaces physical job tickets. It visually tracks the progress of each order and systematically allocates labor and machinery to ensure fulfillment deadlines are met.")
add_definition_item("Printing Services", "Refers to the customizable products and services offered by the establishment (e.g., hardbound binding, document printing, customized merchandise) which are dynamically defined and categorized by the administrator within the system.")
add_definition_item("Production Staff", "Refers to the shop employees who utilize the job scheduling module to view their assigned tasks, update the production progress of specific orders, and initiate the digital deduction of consumed raw materials.")
add_definition_item("Reorder Point", "Refers to the specific, data-driven stock level threshold calculated by the system which, when reached, immediately notifies the administrator that new raw materials must be purchased from suppliers.")

doc.save(docx_path)
print("SUCCESS_DOCX:", docx_path)

# Markdown version
md_content = """# INTEGRATED DYNAMIC ORDER, JOB SCHEDULING, AND INVENTORY MANAGEMENT SYSTEM FOR PRINTING SERVICES WITH AUTOMATED REPLENISHMENT

**KAYE P. COMISSION**  
**CHRISTIAN JAMES J. PEREZ**  

**BACHELOR OF SCIENCE IN INFORMATION SYSTEMS**  
**AUGUST 2026**  

---

# INTRODUCTION

The printing and binding industry is a fast-paced and highly dynamic sector that caters to a diverse range of client needs, from simple document printing and seasonal thesis binding to large-scale marketing materials. As educational institutions, corporate entities, and local businesses continuously generate demand, printing service providers are heavily challenged to maintain operational efficiency. Managing custom orders requires precision not only in recording customer specifications but also in allocating resources, tracking production progress, and ensuring that raw materials are always available. However, many local printing shops still rely on traditional, manual methods or disjointed software that cannot adapt to the varying workflows of different printing services.

Many local printing establishments experience significant operational bottlenecks, especially during peak seasons such as the end of academic semesters or local events. Their current manual setups struggle to seamlessly bridge the gap between order intake, production scheduling, and inventory tracking. When a sudden influx of orders occurs, tracking which employee or machine is assigned to a specific job becomes chaotic. Furthermore, the lack of real-time inventory monitoring often leads to unexpected stockouts of crucial materials—such as specific paper types, ink, and binding covers—resulting in delayed orders, loss of potential revenue, and decreased customer satisfaction.

To address these interconnected challenges, there is a vital need for a system that is not only automated but also highly adaptable to the specific services a shop offers. Traditional management systems are often rigidly hardcoded, preventing business owners from easily adding new product variations or pricing formulas as their business scales. Furthermore, without a data-driven approach to inventory, business owners are forced to guess their restocking schedules based on intuition rather than actual material consumption.

This project aims to design and develop an Integrated Dynamic Order, Job Scheduling, and Inventory Management System for Printing Services with Automated Replenishment. The core innovation of this system lies in its configurable architecture, which acts as a dynamic platform allowing administrators to define, add, or modify their own printing services, pricing variables, and material requirements without altering the source code. To optimize production, a dedicated Job Scheduling module will visually track the progress of each order and allocate labor and machinery efficiently. Finally, to eliminate the problem of material shortages, the system features an Automated Replenishment module. Instead of relying on long-term predictive models that require years of historical data, this module calculates the real-time "burn rate" of raw materials based on active sales volume and automatically alerts the management when stocks reach a critical reorder point.

By integrating these features into a single, cohesive web-based platform, this system intends to modernize the operational workflow of printing establishments, ensuring that orders are fulfilled on time, production is systematically scheduled, and inventory is optimally maintained.

---

## Significance of the Project

The realization of this project will profoundly benefit various stakeholders within the printing and binding industry by transforming disjointed manual processes into a streamlined, automated workflow.

* **Business Owners and Management –** They will gain a dynamic and highly adaptable platform that allows them to effortlessly configure, add, or modify printing services and pricing structures without needing complex source code alterations. Furthermore, the automated inventory replenishment module will eliminate the guesswork in restocking, calculating real-time material burn rates to prevent costly stockouts and maximize potential revenue, especially during peak seasons.
* **Employees and Production Staff –** The system will significantly reduce the chaos associated with sudden influxes of orders. The dedicated job scheduling module will provide a clear, visual queue of assigned tasks and deadlines, preventing overlapping duties and allowing staff to focus efficiently on production rather than manual coordination.
* **Customers and Clients –** They will experience highly reliable service. With optimized order tracking, systematic job delegation, and guaranteed raw material availability, clients will benefit from faster turnaround times, strict adherence to deadlines, and an overall enhanced customer experience.
* **Future Researchers and Developers –** This study will serve as a substantial foundation for future researchers. The integration of a dynamic, configurable service architecture combined with a short-term, data-driven material replenishment logic provides a modern software framework. Future IT and Information Systems scholars can utilize this project as a baseline reference for designing scalable, industry-specific management systems that do not rely on long-term, data-heavy predictive machine learning models.

---

## Statement of the Problem

This study addresses the overarching problem of how to design and develop an Integrated Dynamic Order, Job Scheduling, and Inventory Management System for Printing Services with Automated Replenishment to effectively resolve the manual operational bottlenecks experienced by printing establishments. Currently, the reliance on disjointed, traditional processes creates significant frustrations for all primary stakeholders. For customers, the absence of real-time order tracking leads to the inconvenience of repeatedly visiting or contacting the establishment just to check the status of their print jobs, often resulting in dissatisfaction when orders are unexpectedly delayed. For production staff and employees, depending on manual, paper-based queuing systems creates a chaotic workflow where physical job tickets are easily misplaced, miscommunicated, or improperly prioritized, leading to overlapping duties and unfulfilled deadlines. Meanwhile, business owners and administrators are hindered by rigid software that cannot adapt to new services, and they are forced to rely on pure guesswork and intuition to manage inventory, which inevitably leads to unexpected material stockouts and lost revenue during peak seasons.

To resolve these interconnected issues, the study must explicitly answer how to formulate a dynamic order management architecture that allows administrators to seamlessly configure services and pricing on the fly, eliminating rigid system constraints. Furthermore, it seeks to determine how to develop a job scheduling module that digitizes the production queue to visually track order progress and allocate labor, thereby eliminating lost paperwork and staff confusion. Concurrently, it addresses how to implement an automated inventory replenishment module that calculates real-time material burn rates to replace manual inventory guesswork with data-driven reorder alerts. Finally, the study must answer how to integrate these targeted solutions into a single cohesive platform, and how to evaluate its overall acceptability and performance utilizing the ISO 25010 standard for functional suitability, usability, performance efficiency, and security.

---

## Objectives of the Project

The general objective of the project is to design and develop an Integrated Dynamic Order, Job Scheduling, and Inventory Management System for Printing Services with Automated Replenishment to modernize and automate the operational workflows of printing establishments, from customizable order intake and production tracking to data-driven inventory maintenance.

Specifically, this study aims to:

1. formulate a dynamic order management architecture that enables administrators to configure, add, or modify printing services, pricing variables, and material requirements without altering the source code;
2. develop a job scheduling module that digitizes the production queue, visually tracks order progress, and systematically allocates labor and machinery to ensure timely fulfillment;
3. implement an automated inventory replenishment module that calculates the real-time burn rate of materials based on active sales volume and triggers alerts at critical reorder points;
4. integrate the dynamic ordering, job scheduling, and inventory replenishment modules into a single, cohesive web-based platform; and
5. evaluate the overall acceptability and performance of the developed system using the ISO 25010 software quality standard, specifically focusing on functional suitability, usability, performance efficiency, and security.

---

## Expected Outputs of the Project

The expected outputs of the project will include the following:

1. A fully functional web-based platform that integrates dynamic order configuration, job scheduling, and inventory replenishment tailored specifically for printing and binding service operations.
2. A dynamic order management module that grants business administrators the flexibility to seamlessly create, customize, and modify various printing services and pricing structures without needing backend code alterations.
3. A customer-facing portal that provides an intuitive interface for clients to easily submit their specific printing requirements and track the real-time progress of their orders, eliminating the need for frequent physical follow-ups.
4. A digital job scheduling module that equips production staff with a visual, organized queue of pending and ongoing tasks, allowing for the efficient allocation of labor and machinery to meet strict fulfillment deadlines.
5. An automated inventory replenishment module that continuously tracks material consumption and calculates real-time burn rates based on active sales, providing management with data-driven alerts before essential printing supplies reach critical depletion levels.
6. A comprehensive system documentation and evaluation report detailing the software's performance, functional suitability, usability, and security based on the ISO 25010 quality standards.

---

## Limitations of the Project

The scope of this project is strictly bounded to the order configuration, job scheduling, and inventory replenishment of printing services. The system is not designed to function as a full-scale Enterprise Resource Planning (ERP) software; thus, it completely excludes comprehensive financial accounting, third-party logistics or courier integration, and Human Resource modules such as payroll and biometric employee attendance. Furthermore, the system relies on manual payment verification (e.g., uploading of GCash receipts) and does not integrate with automated bank APIs or live payment gateways. Lastly, the software is purely web-based and does not feature Internet of Things (IoT) capabilities; it will not directly interface with physical printing or binding machines to automatically read hardware statuses or ink levels.

### People
The system is designed to facilitate seamless interaction among three primary users:
* **Administrators / Business Owners:** They have full access to the dynamic configuration of the system. They are responsible for setting up the types of printing services offered, managing pricing structures, monitoring the real-time inventory burn rate, and verifying customer payments.
* **Production Staff / Employees:** They utilize the job scheduling module to view their assigned tasks, update the production progress of each order, and systematically deduct raw materials from the digital inventory as jobs are completed.
* **Customers / Clients:** They access the customer-facing web portal to easily browse offered printing services, submit specific order requirements, upload necessary files, and track the real-time status of their print jobs without needing physical follow-ups.

### Data
The system will process, store, and manage essential operational data to ensure smooth transactions. This includes customer profiles and contact details, transaction histories, and uploaded print files. On the administrative side, the database will handle the dynamic service configurations (custom pricing and variations), real-time job order statuses, and automated inventory records, specifically the current stock levels, material burn rates, and reorder point thresholds.

### Process
The system automates the core operational workflow of the printing establishment. The process begins with dynamic order intake, where customers select services configured by the admin. It proceeds to manual payment verification, followed by the automated routing of the confirmed order into the digital job scheduling queue. During production, the system systematically allocates tasks to specific staff members. As production progresses, the system processes inventory deductions based on material consumption, ultimately triggering automated restocking alerts when supplies reach critical levels.

### Hardware
To effectively run and access the web-based system, users only require standard web-capable devices such as desktop computers, laptops, tablets, or smartphones. A stable internet connection is required for all actors (administrators, staff, and customers) to access the cloud-hosted platform and synchronize data in real-time.

---

## Operational Definition of Terms

To ensure clarity and establish a shared understanding of the concepts discussed throughout this study, the following terms are operationally defined. These definitions reflect how each technical and business-specific term is explicitly utilized within the context of the developed system's dynamic architecture, modules, and workflows, rather than their general dictionary meanings:

* **Administrator –** Refers to the business owner or authorized management personnel who has full access to the system. They are responsible for dynamically configuring printing services, managing pricing variables, monitoring the inventory burn rate, and verifying customer payments.
* **Automated Replenishment –** Refers to the specific system module that eliminates manual stock guesswork by calculating real-time material consumption and automatically triggering alerts when supplies reach a critical level.
* **Burn Rate –** In this study, it refers to the calculated velocity or speed at which specific raw printing materials (e.g., paper, ink, binding covers) are consumed based on active, short-term sales volume.
* **Customer –** Refers to the clients or end-users who utilize the customer-facing web portal to browse offered printing services, submit specific order requirements, upload files, and visually track the real-time status of their print jobs.
* **Dynamic Configuration –** Refers to the core architectural feature of the system that allows administrators to seamlessly create, add, or modify printing services and pricing structures through the user interface without requiring continuous modifications to the backend source code.
* **Inventory Management –** Refers to the digital tracking system used to monitor the availability, quantity, and depletion of raw materials required for production, which automatically updates as production staff complete their assigned tasks.
* **Job Scheduling –** Refers to the digital production queue that replaces physical job tickets. It visually tracks the progress of each order and systematically allocates labor and machinery to ensure fulfillment deadlines are met.
* **Printing Services –** Refers to the customizable products and services offered by the establishment (e.g., hardbound binding, document printing, customized merchandise) which are dynamically defined and categorized by the administrator within the system.
* **Production Staff –** Refers to the shop employees who utilize the job scheduling module to view their assigned tasks, update the production progress of specific orders, and initiate the digital deduction of consumed raw materials.
* **Reorder Point –** Refers to the specific, data-driven stock level threshold calculated by the system which, when reached, immediately notifies the administrator that new raw materials must be purchased from suppliers.
"""

with open(md_path, "w") as f:
    f.write(md_content)

print("SUCCESS_MD:", md_path)
